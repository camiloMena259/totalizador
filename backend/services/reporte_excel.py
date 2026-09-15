"""
Genera un Excel del reporte totalizado, con los mismos colores del frontend.

Hojas:
  1. Banco
  2. Resumen general
  3. Detalle
  4. Movimientos
"""

import re
from datetime import datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from models.schemas import TotalizacionResponse


# Colores del frontend (theme.css + tablas.css)
COLOR_SUPERFICIE = "FFFFFF"
COLOR_BORDE = "D3D7E0"
COLOR_TEXTO = "1D2433"
COLOR_TEXTO_SUAVE = "6B7280"
COLOR_ACENTO = "35606B"
COLOR_ACENTO_SUAVE = "E7EFEF"
COLOR_DEBITO_FONDO = "FCEBEA"
COLOR_DEBITO_TEXTO = "B3423A"
COLOR_CREDITO_FONDO = "E9F6EE"
COLOR_CREDITO_TEXTO = "2F7A4F"
COLOR_FILA_TOTAL = "FAFBFC"
COLOR_BLANCO = "FFFFFF"

FORMATO_MONEDA = '#,##0.00'

NOMBRE_ARCHIVO_DEFAULT = "reporte-extracto"


def sanitizar_nombre_archivo(nombre: str) -> str:
    """Quita caracteres peligrosos y asegura la extensión .xlsx."""
    limpio = (nombre or "").strip()
    limpio = re.sub(r'[\\/:*?"<>|]+', "", limpio)
    limpio = re.sub(r"\s+", " ", limpio).strip(" .")
    if not limpio:
        limpio = NOMBRE_ARCHIVO_DEFAULT
    if not limpio.lower().endswith(".xlsx"):
        limpio = f"{limpio}.xlsx"
    return limpio


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _borde() -> Border:
    lado = Side(style="thin", color=COLOR_BORDE)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _fuente(color: str = COLOR_TEXTO, negrita: bool = False, tamano: int = 11) -> Font:
    return Font(name="Calibri", color=color, bold=negrita, size=tamano)


def _alinear(horizontal: str = "left", vertical: str = "center") -> Alignment:
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=True)


def _pintar_encabezado(celda, texto: str) -> None:
    celda.value = texto
    celda.fill = _fill(COLOR_ACENTO)
    celda.font = _fuente(COLOR_BLANCO, negrita=True, tamano=11)
    celda.alignment = _alinear("center")
    celda.border = _borde()


def _celda_base(celda, valor, *, negrita: bool = False, centro: bool = False, derecha: bool = False) -> None:
    celda.value = valor
    celda.fill = _fill(COLOR_SUPERFICIE)
    celda.font = _fuente(COLOR_TEXTO, negrita=negrita)
    horizontal = "right" if derecha else ("center" if centro else "left")
    celda.alignment = _alinear(horizontal)
    celda.border = _borde()


def _celda_moneda(celda, valor: Optional[float], tipo: Optional[str] = None) -> None:
    """tipo: 'debito' | 'credito' | None. Cero/None se muestran como —."""
    if valor is None or valor == 0:
        _celda_base(celda, "—", centro=True)
        celda.font = _fuente(COLOR_TEXTO_SUAVE)
        return

    celda.value = float(valor)
    celda.number_format = FORMATO_MONEDA
    celda.alignment = _alinear("right")
    celda.border = _borde()

    if tipo == "debito":
        celda.fill = _fill(COLOR_DEBITO_FONDO)
        celda.font = _fuente(COLOR_DEBITO_TEXTO, negrita=True)
    elif tipo == "credito":
        celda.fill = _fill(COLOR_CREDITO_FONDO)
        celda.font = _fuente(COLOR_CREDITO_TEXTO, negrita=True)
    else:
        celda.fill = _fill(COLOR_SUPERFICIE)
        celda.font = _fuente(COLOR_TEXTO, negrita=True)


def _celda_etiqueta(celda, texto: str) -> None:
    celda.value = texto
    celda.fill = _fill(COLOR_ACENTO_SUAVE)
    celda.font = _fuente(COLOR_ACENTO, negrita=True)
    celda.alignment = _alinear("center")
    celda.border = _borde()


def _ancho_columnas(hoja, anchos: dict) -> None:
    for letra, ancho in anchos.items():
        hoja.column_dimensions[letra].width = ancho


def _titulo_hoja(hoja, texto: str, columnas: int) -> None:
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columnas)
    celda = hoja.cell(1, 1, texto)
    celda.fill = _fill(COLOR_ACENTO)
    celda.font = _fuente(COLOR_BLANCO, negrita=True, tamano=14)
    celda.alignment = _alinear("left")
    hoja.row_dimensions[1].height = 28


def _escribir_banco(hoja, banco: str, nombre_archivo: str, resultado: TotalizacionResponse) -> None:
    _titulo_hoja(hoja, "Totalizador de extractos", 2)

    filas = [
        ("Banco", banco),
        ("Archivo", nombre_archivo),
        ("Fecha de generación", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Cantidad de movimientos", len(resultado.movimientos)),
        ("Cantidad de etiquetas", len(resultado.resumen_por_etiqueta)),
    ]

    _pintar_encabezado(hoja.cell(3, 1), "Campo")
    _pintar_encabezado(hoja.cell(3, 2), "Valor")

    for i, (campo, valor) in enumerate(filas, start=4):
        _celda_base(hoja.cell(i, 1), campo, negrita=True)
        hoja.cell(i, 1).font = _fuente(COLOR_TEXTO_SUAVE, negrita=True)
        _celda_base(hoja.cell(i, 2), valor)

    _ancho_columnas(hoja, {"A": 28, "B": 48})
    hoja.sheet_view.showGridLines = False


def _escribir_resumen_general(hoja, resultado: TotalizacionResponse) -> None:
    _titulo_hoja(hoja, "Resumen general", 2)
    _pintar_encabezado(hoja.cell(3, 1), "Indicador")
    _pintar_encabezado(hoja.cell(3, 2), "Valor")

    indicadores = [
        ("Gravamen (4x1000)", resultado.total_gravamen, None),
        ("Intereses", resultado.total_intereses, None),
        ("Total débitos", resultado.total_debitos, "debito"),
        ("Total créditos", resultado.total_creditos, "credito"),
    ]

    for i, (etiqueta, valor, tipo) in enumerate(indicadores, start=4):
        _celda_base(hoja.cell(i, 1), etiqueta)
        _celda_moneda(hoja.cell(i, 2), valor, tipo)
        hoja.row_dimensions[i].height = 22

    _ancho_columnas(hoja, {"A": 28, "B": 22})
    hoja.sheet_view.showGridLines = False


def _escribir_detalle(hoja, resultado: TotalizacionResponse) -> None:
    _titulo_hoja(hoja, "Detalle — resumen por etiqueta", 4)

    encabezados = ["Etiqueta", "Cantidad", "Total débitos", "Total créditos"]
    for col, texto in enumerate(encabezados, start=1):
        _pintar_encabezado(hoja.cell(3, col), texto)

    total_cantidad = 0
    total_debitos = 0.0
    total_creditos = 0.0

    for i, fila in enumerate(resultado.resumen_por_etiqueta, start=4):
        _celda_etiqueta(hoja.cell(i, 1), fila.etiqueta)
        _celda_base(hoja.cell(i, 2), fila.cantidad, derecha=True)
        _celda_moneda(hoja.cell(i, 3), fila.total_debitos, "debito")
        _celda_moneda(hoja.cell(i, 4), fila.total_creditos, "credito")
        total_cantidad += fila.cantidad
        total_debitos += fila.total_debitos
        total_creditos += fila.total_creditos

    fila_total = 4 + len(resultado.resumen_por_etiqueta)
    for col in range(1, 5):
        celda = hoja.cell(fila_total, col)
        celda.fill = _fill(COLOR_FILA_TOTAL)
        celda.font = _fuente(COLOR_TEXTO, negrita=True)
        celda.border = _borde()
        celda.alignment = _alinear("right" if col > 1 else "left")

    hoja.cell(fila_total, 1).value = "Total"
    hoja.cell(fila_total, 2).value = total_cantidad
    hoja.cell(fila_total, 3).value = total_debitos
    hoja.cell(fila_total, 3).number_format = FORMATO_MONEDA
    hoja.cell(fila_total, 4).value = total_creditos
    hoja.cell(fila_total, 4).number_format = FORMATO_MONEDA

    _ancho_columnas(hoja, {"A": 28, "B": 14, "C": 20, "D": 20})
    hoja.auto_filter.ref = f"A3:D{max(3, fila_total - 1)}"
    hoja.freeze_panes = "A4"
    hoja.sheet_view.showGridLines = False


def _escribir_movimientos(hoja, resultado: TotalizacionResponse) -> None:
    _titulo_hoja(hoja, "Movimientos", 6)

    encabezados = ["Día", "Concepto", "Etiqueta", "Débito", "Crédito", "Saldo"]
    for col, texto in enumerate(encabezados, start=1):
        _pintar_encabezado(hoja.cell(3, col), texto)

    for i, mov in enumerate(resultado.movimientos, start=4):
        _celda_base(hoja.cell(i, 1), mov.dia, centro=True)
        _celda_base(hoja.cell(i, 2), mov.concepto)
        _celda_etiqueta(hoja.cell(i, 3), mov.etiqueta)
        _celda_moneda(hoja.cell(i, 4), mov.debito, "debito")
        _celda_moneda(hoja.cell(i, 5), mov.credito, "credito")
        if mov.saldo is None:
            _celda_base(hoja.cell(i, 6), "")
        else:
            _celda_moneda(hoja.cell(i, 6), mov.saldo)

    ultima = 3 + len(resultado.movimientos)
    _ancho_columnas(hoja, {"A": 12, "B": 55, "C": 22, "D": 16, "E": 16, "F": 16})
    if resultado.movimientos:
        hoja.auto_filter.ref = f"A3:F{ultima}"
    hoja.freeze_panes = "A4"
    hoja.sheet_view.showGridLines = False


def generar_excel_reporte(
    resultado: TotalizacionResponse,
    banco: str,
    nombre_archivo: str,
) -> bytes:
    libro = Workbook()

    hoja_banco = libro.active
    hoja_banco.title = "Banco"
    _escribir_banco(hoja_banco, banco, nombre_archivo, resultado)

    hoja_resumen = libro.create_sheet("Resumen general")
    _escribir_resumen_general(hoja_resumen, resultado)

    hoja_detalle = libro.create_sheet("Detalle")
    _escribir_detalle(hoja_detalle, resultado)

    hoja_movimientos = libro.create_sheet("Movimientos")
    _escribir_movimientos(hoja_movimientos, resultado)

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
